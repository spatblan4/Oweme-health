from pathlib import Path
from types import SimpleNamespace

from worker.extract.pdf import extract_pdf_text
from worker.extract.tables import extract_tables


def test_extract_pdf_text_runs_external_script_and_reads_output(tmp_path: Path, monkeypatch):
    pdf_path = tmp_path / "sample.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    script_path = tmp_path / "pdf_to_txt.py"
    script_path.write_text("#!/usr/bin/env python3\n")
    output_text = "page text"

    def fake_run(args, capture_output, text):
        Path(args[-1]).write_text(output_text, encoding="utf-8")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr("worker.extract.pdf.subprocess.run", fake_run)

    text = extract_pdf_text(
        pdf_path,
        deps={
            "script_path": script_path,
            "python_executable": "python3",
        },
    )

    assert text == output_text


def test_extract_tables_parses_pdf_text_into_claim_and_payment_rows(monkeypatch):
    sample_text = """===== Page 1 =====
Leibovsky, Vladimir
Thursday, June 25, 2026
Provider: Leibovsky, Vladimir
Service Date : 06/25/2026
Facility : LAiMA OBGYN INC

===== Page 2 =====
LAiMA OBGYN INC
RECEIPT OF PAYMENT
Date: 06/25/2026
Received From: Chao, Chongchong
Amount: 61.29
Payment Type: Credit Card
"""

    monkeypatch.setattr("worker.extract.tables.extract_pdf_text", lambda path, deps=None: sample_text)

    rows = extract_tables(Path("/tmp/sample.pdf"))

    assert rows[0]["provider"] == "Leibovsky, Vladimir"
    assert rows[0]["service_date"] == "06/25/2026"
    assert rows[1]["merchant"] == "LAiMA OBGYN INC"
    assert rows[1]["amount"] == "61.29"


def test_extract_tables_parses_csv_rows(tmp_path: Path):
    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("provider,service_date\nStone Creek Village Dentistry,2026-07-04\n", encoding="utf-8")

    rows = extract_tables(csv_path)

    assert rows == [
        {
            "provider": "Stone Creek Village Dentistry",
            "service_date": "2026-07-04",
        }
    ]


def test_extract_tables_uses_first_non_empty_xlsx_row_as_header(tmp_path: Path):
    from openpyxl import Workbook

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([])
    sheet.append(["PROVIDER NAME", "SERVICE BEGIN DATE", "PATIENT RESPONSIBILITY"])
    sheet.append(["JAMES KIM", "05/13/2026", 2.37])
    workbook.save(xlsx_path)

    rows = extract_tables(xlsx_path)

    assert rows == [
        {
            "PROVIDER NAME": "JAMES KIM",
            "SERVICE BEGIN DATE": "05/13/2026",
            "PATIENT RESPONSIBILITY": 2.37,
        }
    ]
